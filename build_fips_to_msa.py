#!/usr/bin/env python3
r"""
build_fips_to_msa.py
--------------------
Rebuilds fips_to_msa.json — the county-to-metro label — from the Census Bureau's
Core Based Statistical Area delineation file.

WHAT IT IS FOR
    Results show the metro a county sits in: "Wood County, OH
    (Toledo, OH MSA)". A site selector thinks in metros, not counties, so the
    label is how a result becomes recognisable. It is presentation only - nothing
    scores on it - but a missing or stale label makes a correct match look wrong.

SOURCE: Census CBSA delineation files, "List 1".
    https://www2.census.gov/programs-surveys/metro-micro/geographies
        /reference-files/<year>/delineation-files/list1_<year>.xlsx
    Published when OMB revises the delineations - roughly every ten years with
    smaller updates between - so this changes rarely and matters when it does.
    The newest year is probed rather than pinned.

METROPOLITAN ONLY, deliberately
    The file carries both Metropolitan (1,252 counties, 393 areas) and
    Micropolitan (665 counties) statistical areas. Only Metropolitan is kept,
    which is what the file this replaces did - 1,237 counties across 388 titles,
    on an older vintage. A micropolitan area is a town of 10,000-50,000; calling
    that a "metro" in a result would overstate the market.

    `--include-micro` adds them if you want the wider coverage, at the cost of
    the label meaning something weaker.

CONNECTICUT
    The 2023 delineation is on Connecticut's planning regions, which is the
    geography the rest of the model now uses after migrate_ct_geography.py. The
    file this replaces was on the old counties, so Connecticut's metro labels
    were among the things silently lost - the scorer looked up 09110 and the
    label table only had 09001.

USAGE
    python build_fips_to_msa.py
    python build_fips_to_msa.py --year 2023
    python build_fips_to_msa.py --compare      # report the delta, write nothing
"""
import argparse, io, json, os, ssl, sys, urllib.request
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fips_to_msa.json")
META = os.path.join(HERE, "fips_to_msa_meta.json")

URL = ("https://www2.census.gov/programs-surveys/metro-micro/geographies"
       "/reference-files/%d/delineation-files/list1_%d.xlsx")

METRO = "Metropolitan Statistical Area"


def ssl_context():
    for k in ("SSL_CERT_FILE", "REQUESTS_CA_BUNDLE"):
        p = os.environ.get(k)
        if p and os.path.exists(p):
            try:
                return ssl.create_default_context(cafile=p)
            except Exception:
                pass
    return ssl.create_default_context()


def fetch(year, ctx):
    req = urllib.request.Request(URL % (year, year),
                                 headers={"User-Agent": "FastLocations/1.0"})
    with urllib.request.urlopen(req, timeout=180, context=ctx) as r:
        return r.read()


def latest_year(ctx, newest=None):
    """Newest published delineation, probed downward. OMB revises these on its
    own schedule, so pinning a year means never picking up a revision."""
    top = newest or date.today().year
    for y in range(top, top - 12, -1):
        try:
            fetch(y, ctx)
            return y
        except Exception:
            continue
    return None


def main():
    ap = argparse.ArgumentParser(description="Rebuild fips_to_msa.json")
    ap.add_argument("--year", default="latest")
    ap.add_argument("--include-micro", action="store_true",
                    help="also label micropolitan areas (10,000-50,000 towns)")
    ap.add_argument("--compare", action="store_true")
    a = ap.parse_args()
    ctx = ssl_context()

    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")

    year = latest_year(ctx) if str(a.year).lower() == "latest" else int(a.year)
    if not year:
        sys.exit("could not reach the Census delineation files")
    print("Census CBSA delineation %d" % year)

    wb = openpyxl.load_workbook(io.BytesIO(fetch(year, ctx)), read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))

    # The sheet opens with two title lines; find the header by its own content
    # rather than assuming a row number, which a reformat would break silently.
    hdr_i = next((i for i, r in enumerate(rows[:10])
                  if r and "CBSA Title" in [str(c) for c in r]), None)
    if hdr_i is None:
        sys.exit("could not find the header row - has the sheet been reformatted?")
    hdr = [str(c) if c is not None else "" for c in rows[hdr_i]]
    ix = {h: i for i, h in enumerate(hdr)}
    need = ["CBSA Title", "Metropolitan/Micropolitan Statistical Area",
            "FIPS State Code", "FIPS County Code"]
    missing = [c for c in need if c not in ix]
    if missing:
        sys.exit("delineation file is missing columns: %s" % missing)

    out, micro = {}, 0
    for r in rows[hdr_i + 1:]:
        if not r or not r[0]:
            continue
        kind = str(r[ix["Metropolitan/Micropolitan Statistical Area"]] or "")
        if kind != METRO:
            micro += 1
            if not a.include_micro:
                continue
        st = str(r[ix["FIPS State Code"]] or "").strip()
        co = str(r[ix["FIPS County Code"]] or "").strip()
        title = str(r[ix["CBSA Title"]] or "").strip()
        if not st or not co or not title:
            continue
        out[st.zfill(2) + co.zfill(3)] = title

    old = {}
    if os.path.exists(OUT):
        try:
            with open(OUT, encoding="utf-8") as fh:
                old = json.load(fh)
        except Exception:
            old = {}

    print("  %d counties labelled across %d areas%s"
          % (len(out), len(set(out.values())),
             " (micropolitan included)" if a.include_micro
             else " (%d micropolitan rows skipped)" % micro))
    if old:
        shared = set(out) & set(old)
        renamed = sorted(k for k in shared if out[k] != old[k])
        print("  entries %d -> %d  (%d gained, %d dropped, %d relabelled)"
              % (len(old), len(out), len(set(out) - set(old)),
                 len(set(old) - set(out)), len(renamed)))
        for k in renamed[:4]:
            print("     %s  %s  ->  %s" % (k, old[k], out[k]))
        ct_old = sorted(k for k in old if k.startswith("09"))
        ct_new = sorted(k for k in out if k.startswith("09"))
        if ct_old != ct_new:
            print("  Connecticut: %s -> %s" % (ct_old[:3], ct_new[:3]))

    if a.compare:
        print("--compare: nothing written.")
        return

    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(out, fh, separators=(",", ":"), ensure_ascii=False)
    with open(META, "w", encoding="utf-8") as fh:
        json.dump({"built": date.today().isoformat(),
                   "source": "Census CBSA delineation %d, List 1" % year,
                   "scope": "metropolitan" + (" + micropolitan" if a.include_micro else ""),
                   "counties": len(out),
                   "areas": len(set(out.values()))}, fh, indent=1)
    print("Wrote %s" % OUT)


if __name__ == "__main__":
    main()
